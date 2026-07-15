#!/usr/bin/env python3
"""Build a reviewer-friendly workbook for scale-error annotation.

The input is the raw candidate CSV produced by
``mine_scale_error_candidates_from_eval.py``.  The workbook is meant for human
screening, so it keeps high-recall evidence while separating likely benchmark
items from lower-priority or leakage-prone candidates.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

import pandas as pd
try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    HAS_OPENPYXL = True
except ModuleNotFoundError:
    HAS_OPENPYXL = False


BASE_LIKE_MODELS = {
    "base",
    "qwen3_direct_large",
    "InternVL2.5-1B",
    "scaleprompt_base",
}

SCALE_KEYWORDS = re.compile(
    r"\b(percent|percentage|rate|share|ratio|average|difference|value|sum|max|min|highest|lowest)\b|%",
    flags=re.I,
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-candidates", required=True)
    parser.add_argument("--output-xlsx", required=True)
    parser.add_argument("--output-csv", required=True)
    parser.add_argument("--max-review", type=int, default=500)
    return parser.parse_args()


def uniq(values: Any) -> str:
    vals = [str(v) for v in values if str(v) and str(v) != "nan"]
    return " | ".join(sorted(set(vals)))


def compact_predictions(values: Any, limit: int = 10) -> str:
    vals = []
    for value in values:
        text = str(value).strip()
        if text and text not in vals:
            vals.append(text)
    return " | ".join(vals[:limit])


def auto_decision(row: pd.Series) -> tuple[str, str]:
    priorities = str(row["priorities"])
    factors = str(row["scale_factors"])
    directions = str(row["scale_directions"])
    q = str(row["question"])
    strong_a = "A_scale_100x_single_number" in priorities
    pure_100 = factors in {"100x", "1/100x"}
    has_base = int(row["base_like_occurrences"]) > 0
    many_models = int(row["distinct_models"]) >= 3
    keyword = bool(SCALE_KEYWORDS.search(q))

    if strong_a and pure_100 and has_base and keyword:
        return (
            "yes",
            "High-priority 100x/1-100x single-number error from a base-like model; please verify chart scale.",
        )
    if strong_a and pure_100 and (has_base or many_models):
        return (
            "uncertain",
            "Likely scale relation, but needs chart/image verification before benchmark inclusion.",
        )
    if "B_scale_10x_1000x_single_number" in priorities and has_base:
        return (
            "uncertain",
            "Potential 10x/1000x scale relation from a base-like model; useful for expansion if chart confirms it.",
        )
    if "prediction_too_large" in directions or "prediction_too_small" in directions:
        return (
            "optional",
            "Broad numeric-ratio candidate; review only after high-priority rows.",
        )
    return ("no", "Low-priority candidate.")


def build_table(raw: pd.DataFrame) -> pd.DataFrame:
    groups: list[dict[str, Any]] = []
    for group_id, ((question, gold_answer), grp) in enumerate(
        raw.groupby(["question", "gold_answer"], dropna=False), start=1
    ):
        models = grp["adapter_or_model"].astype(str)
        base_like = grp[models.isin(BASE_LIKE_MODELS)]
        item = {
            "candidate_id": f"SCALE-NAT-{group_id:04d}",
            "suggested_label": "",
            "review_label": "",
            "review_notes": "",
            "question": str(question),
            "gold_answer": str(gold_answer),
            "representative_predictions": compact_predictions(grp["prediction"]),
            "scale_factors": uniq(grp["scale_factor"]),
            "scale_directions": uniq(grp["scale_direction"]),
            "priorities": uniq(grp["priority"]),
            "splits": uniq(grp["split"]),
            "datasets": uniq(grp["dataset"]),
            "occurrences": len(grp),
            "distinct_models": models.nunique(),
            "base_like_occurrences": len(base_like),
            "base_like_models": uniq(base_like["adapter_or_model"]) if len(base_like) else "",
            "all_models": uniq(models),
            "source_files": uniq(grp["source_file"]),
        }
        label, reason = auto_decision(pd.Series(item))
        item["suggested_label"] = label
        item["auto_reason"] = reason
        groups.append(item)

    out = pd.DataFrame(groups)
    label_rank = {"yes": 0, "uncertain": 1, "optional": 2, "no": 3}
    out["label_rank"] = out["suggested_label"].map(label_rank).fillna(9)
    out = out.sort_values(
        ["label_rank", "base_like_occurrences", "occurrences", "distinct_models"],
        ascending=[True, False, False, False],
    ).drop(columns=["label_rank"])
    return out.reset_index(drop=True)


def write_sheet(ws, df: pd.DataFrame, title_fill: str) -> None:
    fill = PatternFill("solid", fgColor=title_fill)
    for col_idx, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
        for col_idx, col in enumerate(df.columns, start=1):
            value = row[col]
            ws.cell(row=row_idx, column=col_idx, value=value)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {
        "candidate_id": 16,
        "suggested_label": 16,
        "review_label": 16,
        "review_notes": 35,
        "question": 70,
        "gold_answer": 16,
        "representative_predictions": 42,
        "scale_factors": 18,
        "scale_directions": 24,
        "priorities": 34,
        "splits": 28,
        "datasets": 18,
        "occurrences": 14,
        "distinct_models": 14,
        "base_like_occurrences": 20,
        "base_like_models": 30,
        "all_models": 60,
        "source_files": 80,
        "auto_reason": 60,
    }
    for idx, col in enumerate(df.columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(col, 18)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    if "review_label" in df.columns:
        col_idx = list(df.columns).index("review_label") + 1
        col_letter = get_column_letter(col_idx)
        dv = DataValidation(type="list", formula1='"yes,no,uncertain,duplicate,not_scale,needs_image"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}{max(2, len(df) + 1)}")


def main() -> None:
    args = parse_args()
    raw = pd.read_csv(args.raw_candidates)
    table = build_table(raw)
    review = table.head(args.max_review).copy()
    yes_uncertain = table[table["suggested_label"].isin(["yes", "uncertain"])].copy()
    optional = table[table["suggested_label"].isin(["optional", "no"])].copy()

    output_xlsx = Path(args.output_xlsx)
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    output_csv = Path(args.output_csv)
    table.to_csv(output_csv, index=False)
    if not HAS_OPENPYXL:
        print("openpyxl is not installed; wrote CSV only:", output_csv)
        print(
            {
                "raw_occurrences": int(len(raw)),
                "distinct_question_gold_candidates": int(len(table)),
                "suggested_counts": table["suggested_label"].value_counts().to_dict(),
                "base_like_candidate_count": int((table["base_like_occurrences"] > 0).sum()),
            }
        )
        return

    wb = Workbook()
    ws_readme = wb.active
    ws_readme.title = "README"
    readme_rows = [
        ("Purpose", "Human review package for natural scale-error candidates mined from model eval CSVs."),
        ("How to label", "Fill review_label. Use yes only if the chart/question/gold/prediction confirms a genuine scale error."),
        ("yes", "Use for benchmark inclusion."),
        ("uncertain", "Needs image or second-pass review."),
        ("duplicate", "Same chart-question already covered by another row."),
        ("not_scale", "Numeric mismatch is not mainly a scale/units error."),
        ("needs_image", "Cannot decide without inspecting the chart."),
        ("Important", "suggested_label is only a machine prefill; reviewer label is authoritative."),
    ]
    for r, (k, v) in enumerate(readme_rows, start=1):
        ws_readme.cell(r, 1, k).font = Font(bold=True)
        ws_readme.cell(r, 2, v).alignment = Alignment(wrap_text=True)
    ws_readme.column_dimensions["A"].width = 18
    ws_readme.column_dimensions["B"].width = 100

    ws_review = wb.create_sheet("review_top")
    write_sheet(ws_review, review, "1F4E79")
    ws_yu = wb.create_sheet("yes_uncertain")
    write_sheet(ws_yu, yes_uncertain, "2F6B45")
    ws_opt = wb.create_sheet("optional")
    write_sheet(ws_opt, optional, "6E5A19")
    ws_all = wb.create_sheet("all_candidates")
    write_sheet(ws_all, table, "444444")

    summary = {
        "raw_occurrences": int(len(raw)),
        "distinct_question_gold_candidates": int(len(table)),
        "suggested_counts": table["suggested_label"].value_counts().to_dict(),
        "base_like_candidate_count": int((table["base_like_occurrences"] > 0).sum()),
    }
    ws_sum = wb.create_sheet("summary")
    for r, (k, v) in enumerate(summary.items(), start=1):
        ws_sum.cell(r, 1, k).font = Font(bold=True)
        ws_sum.cell(r, 2, json.dumps(v, ensure_ascii=False) if isinstance(v, dict) else v)
    ws_sum.column_dimensions["A"].width = 34
    ws_sum.column_dimensions["B"].width = 80

    wb.save(output_xlsx)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print("saved", output_xlsx)
    print("saved", output_csv)


if __name__ == "__main__":
    main()
